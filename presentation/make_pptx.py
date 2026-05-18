import os
from pptx import Presentation
from pptx.util import Inches
from glob import glob

def create_pptx(image_dir, output_file):
    prs = Presentation()
    images = sorted(glob(os.path.join(image_dir, "*.png")))
    for img_path in images:
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)
        slide_width = prs.slide_width
        slide_height = prs.slide_height
        slide.shapes.add_picture(img_path, 0, 0, width=slide_width, height=slide_height)
    prs.save(output_file)
    print(f"Successfully created {output_file}")

if __name__ == "__main__":
    create_pptx("tmp_slides", "thesis_defense.pptx")
